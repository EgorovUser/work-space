"""Разбор содержимого. Лимиты и список форматов — из настроек модуля."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from app.core.logging import logger
from app.core.settings_store import get, get_bool, get_int

OK = "ok"
EMPTY = "empty"
UNSUPPORTED = "unsupported"
ERROR = "error"
SKIPPED = "skipped"


def extract_bytes(filename: str, data: bytes, project_id: int | None) -> tuple[str, str, str]:
    try:
        if not get_bool("documents", "extract.enabled", project_id=project_id):
            return "", SKIPPED, ""
        suffix = Path(filename or "").suffix.lower().lstrip(".")
        allowed = _extensions()
        if suffix not in allowed:
            return "", UNSUPPORTED, ""
        if suffix == "docx":
            text = _docx(data)
        elif suffix == "pdf":
            text = _pdf(data, project_id)
        elif suffix in ("xlsx", "xlsm"):
            text = _xlsx(data, project_id)
        else:
            text = _plain(data)
        text = _clip(text, get_int("documents", "extract.max_chars", project_id=project_id))
        if not text.strip():
            return "", EMPTY, ""
        return text, OK, ""
    except Exception as exc:
        message = _clip_error(exc)
        logger.warning("extract failed name={} error={}", Path(filename or "").name, message)
        return "", ERROR, message


def status_label(status: str) -> str:
    raw = get("documents", "extract.status_labels")
    if isinstance(raw, dict):
        label = raw.get(status)
        if isinstance(label, str) and label.strip():
            return label.strip()
    return status


def _extensions() -> set[str]:
    raw = get("documents", "extract.extensions")
    if not isinstance(raw, list):
        return set()
    return {str(item).lower().lstrip(".") for item in raw if str(item).strip()}


def _docx(data: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _pdf(data: bytes, project_id: int | None) -> str:
    import fitz

    limit = get_int("documents", "extract.pdf.max_pages", project_id=project_id)
    document = fitz.open(stream=data, filetype="pdf")
    try:
        count = min(document.page_count, limit)
        return "\n".join(document.load_page(index).get_text("text") for index in range(count))
    finally:
        document.close()


def _xlsx(data: bytes, project_id: int | None) -> str:
    from openpyxl import load_workbook

    sheets = get_int("documents", "extract.xlsx.max_sheets", project_id=project_id)
    rows = get_int("documents", "extract.xlsx.max_rows", project_id=project_id)
    cols = get_int("documents", "extract.xlsx.max_cols", project_id=project_id)
    book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for index, sheet in enumerate(book.worksheets):
            if index >= sheets:
                break
            parts.append(str(sheet.title or ""))
            for row in sheet.iter_rows(max_row=rows, max_col=cols, values_only=True):
                cells = ["" if value is None else str(value).strip() for value in row]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    parts.append("\t".join(cells))
        return "\n".join(parts)
    finally:
        book.close()


def _plain(data: bytes) -> str:
    raw = get("documents", "extract.encodings")
    encodings = [str(item) for item in raw] if isinstance(raw, list) and raw else ["utf-8"]
    last = "кодировка не подошла"
    for name in encodings:
        try:
            return data.decode(name)
        except (LookupError, UnicodeDecodeError) as exc:
            last = str(exc)
    raise ValueError(last)


def _clip(text: str, limit: int) -> str:
    if limit <= 0:
        return ""
    if len(text) <= limit:
        return text
    return text[:limit]


def _clip_error(exc: Exception) -> str:
    limit = get_int("documents", "extract.error.max_length")
    text = str(exc).strip().replace("\n", " ") or "ошибка разбора"
    if len(text) > limit:
        return text[:limit]
    return text
